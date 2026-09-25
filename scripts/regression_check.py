# -*- coding: utf-8 -*-
# -*- coding: utf-8 -*-
"""培智·双脉教学引擎 —— 全量回归校验（当前引擎版本见 SKILL.md front-matter）

用法：python scripts/regression_check.py      （发布前须 100% PASS）
校验分组：
  ① 教案基准 examples/*.md：BOPPPS 六步 / 每课时时间恒等 / 探究活动 / LOS 成对列 /
     表头行 / 表号章节号 / 占位符 / 红区 / 行为支持卡 / 感官调节 / IEP 累计追踪 / 无障碍说明
  ② 结构化契约 references/output-schema.json 与 examples/*结构化输出样例.json
  ③ 引擎自身 SKILL.md：frontmatter 合规（name/description/version 必填 + 规范白名单）/
     SemVer / 版号一致 / 五条铁律 / 外部引用存在性
  ④ md→JSON 派生一致性（Single Source，含"落盘样例 ≡ 派生结果"防漂移）
  ⑤ Word 成品 md→docx：OOXML 包完整性与元素序列 / 版式契约 / 字节幂等 / 外部读取复校
  ⑥ 交付通道契约：CLI 参数 / PDF 降级 / --check 反篡改（临时目录端到端）
  ⑦ 双轨合规与仓库卫生：技能标准布局（scripts/references/examples）/ 无二进制成品 /
     无运行期报告 / GitHub 治理文件保留且版本同步（README/CHANGELOG）/ 章节号 / 不硬编码项数
输出：%TEMP%/peizhi_shuangmai/regression_report.txt（UTF-8；报告不进仓库）
元断言：同一分节内重复计入的断言必须为 0（防止通过率虚高）
"""
import re
import os
import sys
import json
import glob
import hashlib
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SEP = re.compile(r'^[\s\-:|]+$')

REPORT_DIR = os.path.join(tempfile.gettempdir(), 'peizhi_shuangmai')


def report_path(name):
    """运行期报告一律落系统临时目录（包内不产 *_report.txt）"""
    os.makedirs(REPORT_DIR, exist_ok=True)
    return os.path.join(REPORT_DIR, name)


def blocks(lines):
    """提取 markdown 表格块"""
    out, cur = [], []
    for l in lines:
        if l.strip().startswith('|'):
            cur.append(l.strip())
        else:
            if cur:
                out.append(cur)
                cur = []
    if cur:
        out.append(cur)
    return out


def is_sep(row):
    return bool(SEP.match(row.replace('|', '').strip()))


def cells(row):
    return [c.strip() for c in row.strip().strip('|').split('|')]


def check(path):
    name = os.path.basename(path)
    text = open(path, encoding='utf-8').read()
    lines = text.splitlines()
    r = []

    m = re.search(r'_(\d)课时', name)
    n = int(m.group(1)) if m else 1

    # 1 封面两行标题相邻
    h1 = [i for i, l in enumerate(lines) if l.startswith('# ')]
    r.append(('封面两行标题相邻', len(h1) >= 2 and h1[1] == h1[0] + 1))

    # 2 表格均有表头行（首行非分隔行）
    bs = blocks(lines)
    nohead = [b[0] for b in bs if len(b) >= 2 and is_sep(b[0])]
    r.append(('表格均有表头行', not nohead, '' if not nohead else str(nohead[:2])))

    # 2b 表格内部禁重复分隔行
    dup = []
    for b in bs:
        for row in b[3:]:
            if is_sep(row):
                dup.append(row[:20])
    r.append(('表格内无重复分隔行', not dup, ','.join(dup[:2])))

    # 2c 表格首行不得是数据行（防"缺语义表头"：如分层作业首行写成"第1课时·A"）
    DATA_ROW = re.compile(r'^(第\d+课时|生\d+|\d+′)')
    badhead = [b[0] for b in bs if b and cells(b[0]) and DATA_ROW.match(cells(b[0])[0])]
    r.append(('表格首行非数据行(表头不缺失)', not badhead, str(badhead[:1])))

    # 3 BOPPPS 六步 + 4 时间合计 + 5 探究活动（按五列表逐个校验；表头严格匹配防偏移）
    steps_needed = {'B', 'O', 'P前', 'P参', 'P后', 'S'}
    offhead = [b for b in bs if b and '教学环节' in b[0] and cells(b[0]) and cells(b[0])[0] != '教学环节']
    r.append(('五列表表头严格为"教学环节"', not offhead, str(offhead[:1])))
    proc_tables = [b for b in bs if b and cells(b[0]) and cells(b[0])[0] == '教学环节']
    r.append(('五列表数量=课时数N', len(proc_tables) == n, '%d/%d' % (len(proc_tables), n)))
    for idx, b in enumerate(proc_tables, 1):
        steps, minutes, inquiry = set(), 0, False
        for row in b[2:]:
            c0 = cells(row)[0] if cells(row) else ''
            mm = re.search(r'（([^）·]+)·(\d+)′）', c0)
            if mm:
                steps.add(mm.group(1))
                minutes += int(mm.group(2))
            if '【探究活动】' in row:
                inquiry = True
        r.append(('第%d课时 BOPPPS 六步齐全' % idx, steps_needed <= steps, '缺 ' + str(steps_needed - steps) if not steps_needed <= steps else ''))
        r.append(('第%d课时 时间合计=35′' % idx, minutes == 35, str(minutes)))
        r.append(('第%d课时 含≥1探究活动' % idx, inquiry))

    # 6 LOS 记录表逐课时成对列：1 + 2N + 1
    los = [b for b in bs if b and '起始LOS' in b[0]]
    if los:
        cols = len(cells(los[0][0]))
        r.append(('LOS表逐课时成对列(1+2N+1=%d)' % (2 * n + 2), cols == 2 * n + 2, '实%d列' % cols))
        rows = [c for c in los[0][2:] if cells(c)[0].startswith('生')]
        r.append(('LOS表全生覆盖(12人)', len(rows) == 12, '实%d人' % len(rows)))
    else:
        r.append(('LOS变化记录表存在', False))

    # 7 材料清单 ☐ 勾选
    seg = text.split('材料清单')[1][:600] if '材料清单' in text else ''
    r.append(('材料清单含☐勾选框', '☐' in seg))

    # 8 安全替代表三列含风险
    safe = [b for b in bs if b and '安全替代' in b[0]]
    r.append(('安全替代表三列(含风险列)', bool(safe) and len(cells(safe[0][0])) == 3 and '风险' in safe[0][0]))

    # 9 板书版式为代码块
    r.append(('板书为版式代码块', bool(re.search(r'### 板书与图卡设计\s*\n\s*```', text))))

    # 10 教务归档视图含课标锚点三级（≥2 个“·”分隔）
    arch = text.split('教务归档视图')[1] if '教务归档视图' in text else ''
    anchor = re.search(r'课标依据\s*\|\s*([^|]+)', arch)
    r.append(('归档视图含课标锚点三级', bool(anchor) and anchor.group(1).count('·') >= 2))

    # 11 表号连续
    nums = [int(x) for x in re.findall(r'\*\*表(\d+)', text)]
    r.append(('表号连续无跳号', nums == list(range(1, len(nums) + 1)), str(nums)))

    # 12 章节中文数字编号连续
    cn = re.findall(r'^## ([一二三四五六七八九十]+)、', text, re.M)

    def cn2num(s):
        d = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6, '七': 7, '八': 8, '九': 9}
        if s in d:
            return d[s]
        if s.startswith('十'):
            return 10 + (d.get(s[1:], 0) if len(s) > 1 else 0)
        return 0

    idxs = [cn2num(c) for c in cn]
    r.append(('章节编号连续', idxs == list(range(1, len(idxs) + 1)), '/'.join(cn)))

    # 13 占位符保留
    r.append(('占位符{{}}保留', '{{}}' in text))

    # 14 红区扫描
    hits = []
    if re.search(r'\d{17}[\dXx]', text):
        hits.append('疑似身份证')
    if re.search(r'(?<!\d)1[3-9]\d{9}(?!\d)', text):
        hits.append('疑似手机号')
    if re.search(r'\d{11,}', text):
        hits.append('≥11位数字串')
    r.append(('红区扫描干净', not hits, ','.join(hits)))

    # ===== 2.2.0 新增：跨课时行为干预 =====
    card_sec = ''
    mm = re.search(r'^## [一二三四五六七八九十]+、跨课时行为干预支持卡.*?(?=^## )', text, re.M | re.S)
    if mm:
        card_sec = mm.group(0)
    r.append(('跨课时行为支持卡章节存在', bool(card_sec)))
    keys = ['触发信号', '前因调整', '替代行为', '强化计划', '危机处置', '全员一致要求']
    miss = [k for k in keys if k not in card_sec]
    r.append(('行为支持卡六要素齐全', not miss, '缺' + ','.join(miss) if miss else ''))
    r.append(('强化计划逐课时削弱', '连续强化' in card_sec and ('FR2' in card_sec or '变动强化' in card_sec)))
    r.append(('替代行为标注已先教先练', '≥2次' in card_sec or '2次' in card_sec))
    r.append(('晋级降级阈值含80%判据', '80%' in card_sec and '无参与' in card_sec and '危机' in card_sec))
    taboo = [t for t in ['体罚', '惩罚性隔离', '强制进食', '当众批评']
             if not re.search(r'禁(止)?\s*' + t, card_sec)]
    r.append(('危机处置禁忌齐全(4条)', not taboo, ','.join(taboo)))

    # ===== 2.4.0 新增：五维感官调节前置 =====
    sn_sec = ''
    msn = re.search(r'^### .*感官调节.*?(?=^## |^### (?!.*感官调节))', text, re.M | re.S)
    if msn:
        sn_sec = msn.group(0)
    r.append(('感官调节章节存在', bool(sn_sec)))
    dims = ['听觉', '视觉', '前庭', '口欲', '触觉']
    miss_d = [d for d in dims if d not in sn_sec]
    r.append(('感官五维度齐全', not miss_d, '缺' + ','.join(miss_d) if miss_d else ''))
    sn_tbl = [b for b in bs if b and cells(b[0]) and cells(b[0])[0] == '感官/环境维度']
    r.append(('感官调节表四列(维度/触发/前置/降刺激通道)',
              bool(sn_tbl) and len(cells(sn_tbl[0][0])) == 4, '' if sn_tbl else '缺表'))
    r.append(('降刺激通道非惩罚性(无"隔离"作唯一通道)',
              bool(sn_sec) and '惩罚性' in sn_sec))

    # ===== 2.4.0 新增：IEP 长期目标跨课时累计追踪 =====
    ie_sec = ''
    mie = re.search(r'^### .*IEP.*?(?=^## |^### (?!.*IEP))', text, re.M | re.S)
    if mie:
        ie_sec = mie.group(0)
    r.append(('IEP累计追踪章节存在', bool(ie_sec)))
    ie_tbl = [b for b in bs if b and '年度长期目标' in b[0]]
    r.append(('IEP追踪表含年度长期目标列', bool(ie_tbl)))
    if ie_tbl:
        hdr = cells(ie_tbl[0][0])
        reach = len([h for h in hdr if re.search(r'课时\d+达成', h)])
        r.append(('IEP追踪表逐课时达成列==N', reach == n, '实%d列/N=%d' % (reach, n)))
        rows = [c for c in ie_tbl[0][2:] if cells(c) and cells(c)[0].startswith('生')]
        r.append(('IEP追踪表全生覆盖(12人)', len(rows) == 12, '实%d人' % len(rows)))
        r.append(('IEP累计口径含成功率算式', '累计' in ie_sec and '成功率' in ie_sec and '÷' in ie_sec))

    # ===== 2.2.0 新增：排版与无障碍执行说明 =====
    acc_sec = ''
    ma = re.search(r'^## [一二三四五六七八九十]+、排版与无障碍执行说明.*?(?=^## )', text, re.M | re.S)
    if ma:
        acc_sec = ma.group(0)
    r.append(('排版无障碍执行说明章节存在', bool(acc_sec)))
    miss_a = [k for k in ['12pt', '24pt', '36pt', '7:1', '4.5:1', '不得仅依赖颜色'] if k not in acc_sec]
    r.append(('无障碍基线要素齐全', not miss_a, '缺' + ','.join(miss_a) if miss_a else ''))

    return name, r


def read_engine_version(root):
    """从 SKILL.md frontmatter 读取引擎版本（避免脚本内硬编码）"""
    try:
        text = open(os.path.join(root, 'SKILL.md'), encoding='utf-8').read()
        m = re.search(r'^version: (.+)$', text, re.M)
        return m.group(1).strip() if m else ''
    except Exception:
        return ''


def check_schema(root, ev=''):
    """结构化输出契约与样例校验"""
    r = []
    sp = os.path.join(root, 'references', 'output-schema.json')
    try:
        schema = json.load(open(sp, encoding='utf-8'))
        r.append(('Schema 文件为合法 JSON', True))
    except Exception as e:
        return [('Schema 文件为合法 JSON', False, str(e))]

    top = schema.get('required', [])
    r.append(('Schema 顶层 required 齐全', len(top) >= 8, str(len(top))))
    try:
        step_enum = schema['properties']['lessons']['items']['properties']['timeline']['items']['properties']['step']['enum']
    except KeyError:
        step_enum = []
    r.append(('Schema 含 BOPPPS step 枚举', set(step_enum) == {'B', 'O', 'P前', 'P参', 'P后', 'S'}, str(step_enum)))
    anc = schema.get('properties', {}).get('anchors', {})
    anc_req = anc.get('items', {}).get('required', [])
    sensory = schema.get('properties', {}).get('sensory_regulation', {})
    iep = schema.get('properties', {}).get('iep_tracking', {})
    r.append(('Schema anchors 为多锚点数组', anc.get('type') == 'array' and set(anc_req) == {'板块', '条目', '表述'}))
    sd = sensory.get('items', {}).get('required', [])
    r.append(('Schema 感官调节四字段', set(sd) == {'维度', '触发信号', '前置安排', '降刺激通道'}, str(sd)))
    r.append(('Schema 感官调节≥5维度', sensory.get('minItems') == 5 and
              set(sensory.get('items', {}).get('properties', {}).get('维度', {}).get('enum', [])) ==
              {'听觉', '视觉', '前庭与座位', '口欲与过敏', '触觉与材料'}))
    idf = iep.get('items', {}).get('required', [])
    r.append(('Schema IEP 追踪五字段',
              set(idf) == {'学生代号', '年度长期目标', '本课短期目标', '分课时达成', '累计口径'}, str(idf)))
    r.append(('Schema 顶层含 sensory_regulation/iep_tracking',
              'sensory_regulation' in schema.get('required', []) and 'iep_tracking' in schema.get('required', [])))
    try:
        bc = schema['properties']['support']['properties']['behavior_card']['required']
    except KeyError:
        bc = []
    r.append(('Schema 行为支持卡六要素', len(bc) == 6, str(len(bc))))

    jf = glob.glob(os.path.join(root, 'examples', '*结构化输出样例.json'))
    if not jf:
        return r + [('存在结构化输出样例', False)]
    try:
        data = json.load(open(jf[0], encoding='utf-8'))
        r.append(('JSON 样例为合法 JSON', True))
    except Exception as e:
        return r + [('JSON 样例为合法 JSON', False, str(e))]

    r.append(('JSON schema_version 符合引擎版本', data.get('schema_version') == ev,
              'JSON=%s SKILL=%s' % (data.get('schema_version'), ev)))
    miss_top = [k for k in top if k not in data]
    r.append(('JSON 顶层字段无缺失', not miss_top, '缺' + ','.join(miss_top) if miss_top else ''))
    # SKILL.md §6 声明的顶层字段集合必须与 Schema required 一致（防文档与契约漂移）
    try:
        sm = open(os.path.join(root, 'SKILL.md'), encoding='utf-8').read()
        decl = re.search(r'顶层必填\s*`([^`]+)`', sm)
        names = set(re.findall(r'[A-Za-z_][A-Za-z0-9_]*', decl.group(1))) if decl else set()
        r.append(('SKILL.md §6 顶层字段声明 ≡ Schema required',
                  bool(names) and names == set(top),
                  '多%s 缺%s' % (','.join(sorted(names - set(top))),
                                 ','.join(sorted(set(top) - names)))))
    except Exception as e:
        r.append(('SKILL.md §6 顶层字段声明 ≡ Schema required', False, str(e)[:60]))
    ancs = data.get('anchors', [])
    r.append(('JSON 多锚点且每条三级齐全',
              len(ancs) >= 1 and all(all(a.get(k) for k in ('板块', '条目', '表述')) for a in ancs),
              '%d条' % len(ancs)))

    n = data['meta']['课时数N']
    dur = data['meta']['单课时时长分钟']
    lessons = data['lessons']
    r.append(('JSON lessons 数 == 课时数N', len(lessons) == n, '%d/%d' % (len(lessons), n)))

    needed = {'B', 'O', 'P前', 'P参', 'P后', 'S'}
    steps_ok, time_ok, matrix_ok, inq_ok = True, True, True, True
    detail = []
    for ls in lessons:
        steps = set(x['step'] for x in ls['timeline'])
        s = sum(x['分钟'] for x in ls['timeline'])
        if not needed <= steps:
            steps_ok = False
        if s != dur:
            time_ok = False
        if len(ls['objectives_matrix']) < 9:
            matrix_ok = False
        if len(ls['inquiry_activities']) < 1:
            inq_ok = False
        detail.append('L%d=%d' % (ls['课时序号'], s))
    r.append(('JSON 每课时 BOPPPS 六步齐全', steps_ok))
    r.append(('JSON 每课时分钟合计==%d' % dur, time_ok, ' '.join(detail)))
    r.append(('JSON 每课时目标矩阵≥9格', matrix_ok))
    r.append(('JSON 每课时探究活动≥1', inq_ok))

    los_ok, code_ok = True, True
    for row in data['los_table']:
        recs = row['records']
        if len(recs) != n:
            los_ok = False
        if any(('起始LOS' not in x or '达成LOS' not in x) for x in recs):
            los_ok = False
        if not re.match(r'^生\d+$', row['学生代号']):
            code_ok = False
    r.append(('JSON LOS 逐生逐课时成对', los_ok, '%d人' % len(data['los_table'])))
    r.append(('JSON 学生一律代号(红区合规)', code_ok))

    sup = data['support']
    keys = {'触发信号', '前因调整', '替代行为', '强化计划', '危机处置', '全员一致要求'}
    r.append(('JSON 行为支持卡六要素', keys <= set(sup['behavior_card'])))
    r.append(('JSON 含晋级降级阈值规则', len(sup.get('progression_rules', {})) >= 5))
    r.append(('JSON 泛化三场景齐全', all(data['generalization'].get(k) for k in ('家庭', '学校', '社区'))))
    r.append(('JSON 安全替代表三字段', all({'环节材料', '风险', '安全替代'} <= set(x) for x in data['safety_alternatives'])))
    r.append(('JSON 隐私声明 red_zone_free', data['privacy'].get('red_zone_free') is True))
    # 2.4.0：感官调节与 IEP 追踪
    sdims = [x['维度'] for x in data.get('sensory_regulation', [])]
    r.append(('JSON 感官五维度齐全',
              set(sdims) == {'听觉', '视觉', '前庭与座位', '口欲与过敏', '触觉与材料'}, '/'.join(sdims)))
    r.append(('JSON 感官每条四字段非空',
              all(all(x.get(k) for k in ('触发信号', '前置安排', '降刺激通道'))
                  for x in data.get('sensory_regulation', []))))
    ie_rows = data.get('iep_tracking', [])
    r.append(('JSON IEP 追踪全生覆盖(12人)', len(ie_rows) == 12, '实%d人' % len(ie_rows)))
    r.append(('JSON IEP 追踪逐课时达成列==N',
              all(len(row.get('分课时达成', {})) == n for row in ie_rows)))
    r.append(('JSON IEP 追踪字段非空且代号合规',
              bool(ie_rows) and all(row.get('年度长期目标') and row.get('本课短期目标')
                                    and re.match(r'^生\d+$', row['学生代号']) and row.get('累计口径')
                                    for row in ie_rows)))
    return r


def check_derived(root, ev=''):
    """md → JSON 派生一致性校验（Single Source 硬保证）"""
    r = []
    sys.path.insert(0, os.path.join(root, 'scripts'))
    try:
        import md_to_json as M
        r.append(('派生器 ENGINE_VERSION 与引擎一致', M.ENGINE_VERSION == ev,
                  'M=%s SKILL=%s' % (M.ENGINE_VERSION, ev)))
    except Exception as e:
        return [('可导入 md_to_json', False, str(e))]
    try:
        import md_to_docx as D
        r.append(('Word 生成器版本与引擎一致', D.ENGINE_VERSION == ev,
                  'D=%s SKILL=%s' % (D.ENGINE_VERSION, ev)))
    except Exception as e:
        r.append(('可导入 md_to_docx', False, str(e)))
    files = sorted(glob.glob(os.path.join(root, 'examples', '*.md')))
    for f in files:
        nm = os.path.basename(f)
        try:
            data = M.derive(open(f, encoding='utf-8').read())
        except Exception as e:
            r.append(('[%s] 派生成功' % nm, False, str(e)[:60]))
            continue
        n = data['meta']['课时数N']
        dur = data['meta']['单课时时长分钟']
        tag = '[%s]' % nm
        r.append((tag + ' 派生 lessons 数==N', len(data['lessons']) == n, '%d/%d' % (len(data['lessons']), n)))
        bad = []
        for ls in data['lessons']:
            s = sum(x['分钟'] for x in ls['timeline'])
            if s != dur:
                bad.append('L%d=%d' % (ls['课时序号'], s))
        r.append((tag + ' 派生每课时时间合计==%d' % dur, not bad, ' '.join(bad)))
        r.append((tag + ' 派生 LOS 逐课时成对',
                  len(data['los_table']) >= 1 and all(len(row['records']) == n for row in data['los_table']),
                  '%d人' % len(data['los_table'])))
        r.append((tag + ' 派生目标矩阵≥9格', all(len(ls['objectives_matrix']) >= 9 for ls in data['lessons'])))
        r.append((tag + ' 派生探究活动≥1/课时', all(len(ls['inquiry_activities']) >= 1 for ls in data['lessons'])))
        r.append((tag + ' 派生分层作业三层齐全',
                  all(all(ls['分层作业'].get(k) for k in ('A', 'B', 'C')) for ls in data['lessons'])))
        r.append((tag + ' 派生锚点非空且三级',
                  len(data['anchors']) >= 1 and all(all(a.get(k) for k in ('板块', '条目', '表述')) for a in data['anchors']),
                  '%d条' % len(data['anchors'])))
        r.append((tag + ' 派生行为支持卡六要素', len(data['support']['behavior_card']) >= 6,
                  str(len(data['support']['behavior_card']))))
        r.append((tag + ' 派生泛化三场景', all(data['generalization'].get(k) for k in ('家庭', '学校', '社区'))))
        r.append((tag + ' 派生安全替代表非空', len(data['safety_alternatives']) >= 1))
        r.append((tag + ' 派生材料清单非空', len(data['materials']) >= 1))
        p = os.path.join(root, 'examples', '%s_结构化输出样例.json' % data['meta']['课题'])
        if os.path.exists(p):
            saved = json.load(open(p, encoding='utf-8'))
            same = json.dumps(saved, sort_keys=True, ensure_ascii=False) == json.dumps(data, sort_keys=True, ensure_ascii=False)
            r.append((tag + ' 落盘样例≡派生结果(防漂移)', same, '' if same else '须重跑 scripts/md_to_json.py'))
    return r


def check_skill_md(root):
    """SKILL.md 引擎自身防漂移校验（frontmatter / 版本号 / 铁律 / 新增机制 / 外部引用）"""
    r = []
    p = os.path.join(root, 'SKILL.md')
    text = open(p, encoding='utf-8').read()
    fm = re.match(r'^---\n(.*?)\n---\n', text, re.S)
    if not fm:
        return [('SKILL.md frontmatter 存在', False)]
    head = fm.group(1)
    keys = [l.split(':', 1)[0] for l in head.splitlines() if l and not l.startswith(' ')]
    # —— AgentKit / 豆包 / 千问技能规范（2.5.0 起）：name+description 必填、version 平台解析、
    #    其余顶层字段仅限规范白名单（license/compatibility/metadata/allowed-tools）——
    need = ['name', 'description', 'version']
    miss = [k for k in need if k not in keys]
    r.append(('frontmatter 必填字段齐全(name/description/version)', not miss,
              '缺' + ','.join(miss) if miss else ''))
    allowed = {'name', 'description', 'version', 'license', 'compatibility',
               'metadata', 'allowed-tools'}
    extra = [k for k in keys if k not in allowed]
    r.append(('frontmatter 顶层字段不超出规范白名单', not extra, ','.join(extra)))
    banned = ['display_name', 'display_name_en', 'description_zh', 'description_en',
              'tags', 'requires', 'author']
    left = [k for k in banned if k in keys]
    r.append(('旧版非标顶层字段已移除(防平台解析异常)', not left, ','.join(left)))
    nm = re.search(r'^name: (.+)$', head, re.M)
    nm_v = nm.group(1).strip() if nm else ''
    r.append(('name 合规(小写+连字符,≤64,不含agentkit)',
              bool(re.match(r'^[a-z0-9]+(-[a-z0-9]+)*$', nm_v))
              and len(nm_v) <= 64 and 'agentkit' not in nm_v, nm_v))
    ds = re.search(r'^description: (.+)$', head, re.M)
    ds_v = ds.group(1).strip() if ds else ''
    r.append(('description 非空且≤1024字符', 0 < len(ds_v) <= 1024, '%d字符' % len(ds_v)))
    r.append(('description 不含 XML 标签', not re.search(r'<[A-Za-z/!]', ds_v)))
    ver = re.search(r'^version: (.+)$', head, re.M)
    ok_semver = bool(ver and re.match(r'^\d+\.\d+\.\d+$', ver.group(1).strip()))
    r.append(('SKILL.md version 为 SemVer 三段', ok_semver, ver.group(1) if ver else ''))
    # 版本号多处一致：frontmatter / 正文标题 / JSON 样例 schema_version（脚本与 Schema 由派生组校验）
    v = ver.group(1).strip() if ver else ''
    jf = glob.glob(os.path.join(root, 'examples', '*结构化输出样例.json'))
    sv = ''
    if jf:
        try:
            sv = json.load(open(jf[0], encoding='utf-8')).get('schema_version', '')
        except Exception:
            sv = 'ERR'
    r.append(('JSON 样例版本与引擎一致', sv == v, 'JSON=%s SKILL=%s' % (sv, v)))
    r.append(('SKILL.md 正文标题含版本号', ('V' + v) in text, 'V' + v))
    # 五条铁律不漂移
    laws = ['红区零输入零输出', '学情禁编造', '教材禁杜撰', '课标锚点三级', '时长恒等']
    miss_l = [k for k in laws if k not in text]
    r.append(('五条铁律表述保留', not miss_l, '缺' + ','.join(miss_l) if miss_l else ''))
    # 2.2.0 新增机制保留
    mech = ['跨课时行为干预递进', 'ABC 简录', '连续 2 课时', '80%', '连续强化',
            '危机处置', '跨课时行为支持卡', '无参与(N)']
    miss_m = [k for k in mech if k not in text]
    r.append(('行为干预机制保留', not miss_m, '缺' + ','.join(miss_m) if miss_m else ''))
    acc = ['7:1', '4.5:1', '24pt', '36pt', '不得仅依赖颜色']
    miss_a = [k for k in acc if k not in text]
    r.append(('无障碍基线保留', not miss_a, '缺' + ','.join(miss_a) if miss_a else ''))
    # 外部引用存在性（防改名失联）
    refs = re.findall(r'`((?:scripts|references)/[\w\.\-]+|examples/[\w一-龥\.\-]+\.(?:md|json))`', text)
    broken = [x for x in set(refs) if not os.path.exists(os.path.join(root, x.replace('/', os.sep)))]
    r.append(('SKILL.md 外部引用均存在', not broken, ','.join(broken)))
    return r


def check_docx(root):
    """Word 成品校验（2.4.0）：结构 / 版式契约 / 字节幂等 / 与源稿防漂移 / 外部读取复校"""
    import zipfile
    import xml.dom.minidom as minidom
    r = []
    sys.path.insert(0, os.path.join(root, 'scripts'))
    try:
        import md_to_docx as G
        r.append(('可导入 md_to_docx 生成器', True))
    except Exception as e:
        return [('可导入 md_to_docx 生成器', False, str(e)[:60])]

    for f in sorted(glob.glob(os.path.join(root, 'examples', '*.md'))):
        nm = os.path.basename(f)
        tag = '[%s]' % nm
        text = open(f, encoding='utf-8').read()
        try:
            data = G.build_docx_bytes(text)
        except Exception as e:
            r.append((tag + ' docx 生成成功', False, str(e)[:60]))
            continue
        r.append((tag + ' docx 生成成功', True))

        # ① 字节幂等：同一输入两次生成必须一致
        same_idem = hashlib.sha256(data).hexdigest() == hashlib.sha256(
            G.build_docx_bytes(text)).hexdigest()
        r.append((tag + ' docx 字节级幂等', same_idem))

        # ② OOXML 结构良构 + 版式契约（2.5.0 起 docx 不入库作基准，
        #    "成品≡源稿"端到端校验由交付通道组在临时目录执行）
        try:
            z = zipfile.ZipFile(__import__('io').BytesIO(data))
            doc = z.read('word/document.xml').decode('utf-8')
            sty = z.read('word/styles.xml').decode('utf-8')
            ftr = z.read('word/footer1.xml').decode('utf-8')
            minidom.parseString(doc)
            r.append((tag + ' document.xml 良构', True))
        except Exception as e:
            r.append((tag + ' document.xml 良构', False, str(e)[:60]))
            continue
        first = doc.index('<w:sectPr>')
        cover_sp = doc[first:doc.index('</w:sectPr>', first)]
        body_sp = doc[doc.rindex('<w:sectPr>'):]
        r.append((tag + ' 封面独立成节(sectPr=2)', doc.count('<w:sectPr') == 2,
                  '实%d' % doc.count('<w:sectPr')))
        r.append((tag + ' 封面节无页脚引用(封面不出现页码)', 'footerReference' not in cover_sp))
        r.append((tag + ' 正文节含页脚且正文起始页码=1',
                  'footerReference' in body_sp and 'w:start="1"' in body_sp))
        r.append((tag + ' 页脚为 PAGE/SECTIONPAGES 字段',
                  'PAGE' in ftr and 'SECTIONPAGES' in ftr))
        r.append((tag + ' 页面 A4 与页边距 2.54/3.18cm',
                  'w:w="11906"' in body_sp and 'w:h="16840"' in body_sp
                  and 'w:top="1440"' in body_sp and 'w:left="1803"' in body_sp))
        r.append((tag + ' 默认字体宋体/TimesNewRoman/小四',
                  all(k in sty for k in ('w:eastAsia="宋体"', 'w:ascii="Times New Roman"',
                                         '<w:sz w:val="24"/>'))))
        r.append((tag + ' 一级标题黑体三号加粗',
                  'w:eastAsia="黑体"' in sty and '<w:sz w:val="32"/>' in sty and '<w:b/>' in sty))
        used = set(re.findall(r'<w:pStyle w:val="(Heading\d)"/>', doc))
        r.append((tag + ' 章节标题用 Heading 样式(keepNext 禁标题孤行)',
                  len(used) >= 2 and '<w:keepNext/>' in sty, ','.join(sorted(used))))

        tbls = re.findall(r'<w:tbl>(.*?)</w:tbl>', doc, re.S)
        bad_w, heads, ratio_bad = [], 0, []
        for t in tbls:
            grid = [int(x) for x in re.findall(r'<w:gridCol w:w="(\d+)"', t)]
            row0 = re.search(r'<w:tr>(.*?)</w:tr>', t, re.S).group(1)
            if sum(grid) != 8300:
                bad_w.append(sum(grid))
            if '<w:tblHeader/>' in row0:
                heads += 1
            if len(grid) == 5:
                want = [round(8300 * x / 100.0) for x in (14, 26, 20, 24, 16)]
                if max(abs(a - b) for a, b in zip(grid, want)) > 3:
                    ratio_bad.append(str(grid))
        r.append((tag + ' 表格宽度合计=版心且不溢出', not bad_w, str(bad_w[:2])))
        r.append((tag + ' 表格表头跨页重复', len(tbls) > 0 and heads == len(tbls),
                  '%d/%d' % (heads, len(tbls))))
        r.append((tag + ' 五列表列宽比 14/26/20/24/16', not ratio_bad, ','.join(ratio_bad[:1])))

        # ③b OOXML 包完整性 + 版式细节 + 无障碍底线
        names = z.namelist()
        ct = z.read('[Content_Types].xml').decode('utf-8')
        rels = z.read('_rels/.rels').decode('utf-8')
        r.append((tag + ' 含 docProps/core.xml 且类型与关系齐全',
                  'docProps/core.xml' in names and '/docProps/core.xml' in ct
                  and 'core-properties' in rels))
        core = z.read('docProps/core.xml').decode('utf-8') if 'docProps/core.xml' in names else ''
        leak2 = [k for k in ('双脉', 'WorkBuddy', 'SKILL', 'python', 'md_to_docx', '引擎')
                 if k in core]
        r.append((tag + ' core.xml 零引擎元信息', not leak2, ','.join(leak2)))
        r.append((tag + ' 正文 1.5 倍行距', 'w:line="360" w:lineRule="auto"' in sty))
        r.append((tag + ' 正文段首缩进 2 字符', 'w:firstLineChars="200"' in doc))
        r.append((tag + ' 表格字号≥小五(9pt)', 'w:sz w:val="18"' in doc))
        r.append((tag + ' 表头加粗居中',
                  re.search(r'<w:tblHeader/>.*?<w:b/>', doc, re.S) is not None))
        fills = set(re.findall(r'w:fill="([0-9A-Fa-f]{6})"', doc))

        def neutral_light(h):                      # 中性灰且浅色（禁彩色底纹作唯一编码）
            r_, g_, b_ = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
            return max(r_, g_, b_) - min(r_, g_, b_) <= 6 and min(r_, g_, b_) >= 0xE0
        colored = [h for h in fills if not neutral_light(h)]
        r.append((tag + ' 底纹为中性浅色(色彩不作唯一编码)', not colored,
                  ','.join(sorted(colored))))
        r.append((tag + ' 封面标题二号居中',
                  'w:sz w:val="44"' in doc and 'w:val="center"' in doc))

        # ③c OOXML 元素序列（Word 对顺序敏感，顺序错会触发"内容不可读"修复提示）
        def seq_ok(fragment, tags):
            idx = [fragment.find(t) for t in tags]
            return all(i >= 0 for i in idx) and idx == sorted(idx)
        r.append((tag + ' sectPr 元素序(footerRef→pgSz→pgMar→pgNumType→cols→docGrid)',
                  seq_ok(body_sp, ['<w:footerReference', '<w:pgSz', '<w:pgMar',
                                   '<w:pgNumType', '<w:cols', '<w:docGrid'])))
        tblpr = re.search(r'<w:tblPr>(.*?)</w:tblPr>', doc, re.S)
        r.append((tag + ' tblPr 元素序(tblW→jc→borders→layout→cellMar)',
                  tblpr is not None and seq_ok(tblpr.group(1),
                                               ['<w:tblW', '<w:jc', '<w:tblBorders',
                                                '<w:tblLayout', '<w:tblCellMar'])))
        r.append((tag + ' trPr 元素序(cantSplit→tblHeader)',
                  seq_ok(doc, ['<w:cantSplit/>', '<w:tblHeader/>'])))
        shd_ps = [m.group(0) for m in re.finditer(r'<w:pPr>.*?</w:pPr>', doc, re.S)
                  if 'w:shd' in m.group(0)]
        r.append((tag + ' pPr 中 shd 在 spacing/ind/jc 之前',
                  bool(shd_ps) and all(seq_ok(s, ['<w:shd', '<w:spacing', '<w:ind'])
                                       or '<w:spacing' not in s for s in shd_ps)))

        all_text = re.sub(r'<w:t[^>]*>([^<]*)</w:t>', r'\1', doc + ftr)
        r.append((tag + ' 占位符{{}}在成品中保留', '{{}}' in all_text))
        leak = [k for k in ('双脉教学引擎', 'V2.', 'SKILL.md', 'regression') if k in all_text]
        r.append((tag + ' 成品零引擎元信息', not leak, ','.join(leak)))
        hits = G.red_scan(all_text)
        r.append((tag + ' 成品红区扫描干净', not hits, ','.join(hits)))

        # ④ 外部读取复校（安装 python-docx 时启用）
        try:
            import docx
            d = docx.Document(__import__('io').BytesIO(data))
            txt = '\n'.join(p.text for p in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    txt += '\n' + '|'.join(c.text for c in row.cells)
            ok_par = len(d.paragraphs) > 50 and len(d.tables) > 5
            r.append((tag + ' python-docx 可读且结构完整', ok_par,
                      '段%d/表%d' % (len(d.paragraphs), len(d.tables))))
            r.append((tag + ' 外部读取后占位符与五列表完好',
                      '{{}}' in txt and any('教学环节' in row.cells[0].text
                                            for t2 in d.tables for row in t2.rows
                                            if len(row.cells) == 5)))
        except ImportError:
            pass
    return r


def check_delivery(root):
    """交付通道契约（2.4.0 起）：CLI 参数 / PDF 降级不阻断 / --check 反篡改端到端
    （2.5.0 起 docx 成品不入库，反篡改端到端在系统临时目录执行，不触碰仓库）"""
    r = []
    src = open(os.path.join(root, 'scripts', 'md_to_docx.py'), encoding='utf-8').read()
    r.append(('CLI 支持 --check', "'--check' in sys.argv" in src))
    r.append(('CLI 支持 --pdf', "'--pdf' in sys.argv" in src))
    sys.path.insert(0, os.path.join(root, 'scripts'))
    try:
        import md_to_docx as G
        how = G.export_pdf(os.path.join(root, 'examples', '__none__.docx'),
                           os.path.join(root, 'examples', '__none__.pdf'))
        r.append(('PDF 三通道皆无时降级返回 None(不阻断交付)', how is None, str(how)))
    except Exception as e:
        r.append(('PDF 三通道皆无时降级返回 None(不阻断交付)', False, str(e)[:60]))

    # 在临时目录复刻 `--check` 的判定逻辑（跨进程调用在部分 Windows 控制台下会产生解码噪声；
    # 判定核心是「落盘成品 hash ≡ 源稿派生 hash」，进程内执行等价且更快）
    import md_to_docx as G
    tmp = os.path.join(REPORT_DIR, 'check_tmp')
    os.makedirs(tmp, exist_ok=True)

    def gen_all():
        for f in sorted(glob.glob(os.path.join(root, 'examples', '*.md'))):
            text = open(f, encoding='utf-8').read()
            dest = os.path.join(tmp, G.plan_name(text, os.path.basename(f)[:-3]))
            open(dest, 'wb').write(G.build_docx_bytes(text))

    def run_check():
        states = []
        for f in sorted(glob.glob(os.path.join(root, 'examples', '*.md'))):
            text = open(f, encoding='utf-8').read()
            want = hashlib.sha256(G.build_docx_bytes(text)).hexdigest()
            dest = os.path.join(tmp, G.plan_name(text, os.path.basename(f)[:-3]))
            got = hashlib.sha256(open(dest, 'rb').read()).hexdigest() if os.path.exists(dest) else ''
            states.append(got == want)
        return states

    try:
        gen_all()
        states = run_check()
        r.append(('--check 判定：全部成品 ≡ 源稿(SAME)', all(states) and states,
                  '%d/%d' % (sum(states), len(states))))
        victim = os.path.join(tmp, '认识5_教学设计方案_2课时.docx')
        orig = open(victim, 'rb').read()          # 内存备份（不新建文件，避免触发文件监控进程噪声）
        try:
            with open(victim, 'wb') as fh:
                fh.write(orig[:60] + b'X' + orig[61:])   # 模拟成品被手工改动
            r.append(('--check 判定：能识别被改动的成品(DIFF)',
                      not all(run_check())))
        finally:
            with open(victim, 'wb') as fh:
                fh.write(orig)
        r.append(('--check 判定：还原后重新 SAME', all(run_check())))
    except Exception as e:
        r.append(('--check 判定：全部成品 ≡ 源稿(SAME)', False, str(e)[:60]))
    return r


def check_repo(root, ev):
    """双轨合规与仓库卫生（2.5.0）：技能上传规范（豆包/千问/WorkBuddy）＋ GitHub 仓库规范
    —— 仓库保留治理文件（GitHub 轨），上传 zip 由 scripts/build_package.py 从仓库产出（技能轨）"""
    r = []
    # ① 技能标准布局（scripts/ + references/ + examples/）
    r.append(('scripts/ 四脚本齐全(含打包器)', all(os.path.exists(os.path.join(root, 'scripts', f)) for f in
              ('md_to_docx.py', 'md_to_json.py', 'regression_check.py', 'build_package.py'))))
    r.append(('references/ 三件齐全', all(os.path.exists(os.path.join(root, 'references', f)) for f in
              ('output-schema.json', 'format-baseline.md', 'release-checklist.md'))))
    r.append(('docs/ 自定义目录已移除', not os.path.exists(os.path.join(root, 'docs'))))
    # ② GitHub 轨：治理文件保留且与当前版本同步
    gov = [f for f in ('README.md', 'CHANGELOG.md', 'CONTRIBUTING.md', 'LICENSE', '.gitignore')
           if not os.path.exists(os.path.join(root, f))]
    r.append(('GitHub 治理文件保留(README/CHANGELOG/CONTRIBUTING/LICENSE/.gitignore)',
              not gov, '缺' + ','.join(gov) if gov else ''))
    try:
        rd = open(os.path.join(root, 'README.md'), encoding='utf-8').read()
        r.append(('README 标题含当前版本', ev in rd.split('\n', 1)[0],
                  rd.split('\n', 1)[0][:40]))
        cl = open(os.path.join(root, 'CHANGELOG.md'), encoding='utf-8').read()
        r.append(('CHANGELOG 含当前版本条目', '## [%s]' % ev in cl))
    except Exception as e:
        r.append(('README/CHANGELOG 可读', False, str(e)[:60]))
    # ③ 包体卫生：无二进制成品、无运行期报告、无临时调试件（递归，跳过隐藏目录与 dist）
    bins = glob.glob(os.path.join(root, 'examples', '*.docx')) + \
        glob.glob(os.path.join(root, 'examples', '*.pdf'))
    r.append(('examples/ 无二进制成品(docx/pdf)', not bins,
              ','.join(os.path.basename(b) for b in bins[:3])))
    junk = []
    for dp, dn, fn in os.walk(root):
        dn[:] = [d for d in dn if not d.startswith('.') and d != 'dist']
        for f in fn:
            if f.endswith('_report.txt') or f.startswith('_'):
                junk.append(os.path.relpath(os.path.join(dp, f), root))
    r.append(('仓库无运行期报告与临时调试件', not junk, ','.join(junk[:3])))
    # ④ SKILL.md 体积（渐进式披露：主文件 ≤500 行）
    n_lines = len(open(os.path.join(root, 'SKILL.md'), encoding='utf-8').read().splitlines())
    r.append(('SKILL.md 行数≤500(渐进式披露)', n_lines <= 500, '%d行' % n_lines))
    # ⑤ .gitignore：覆盖缓存/工作目录/打包产物
    gi_p = os.path.join(root, '.gitignore')
    gi = open(gi_p, encoding='utf-8').read() if os.path.exists(gi_p) else ''
    r.append(('.gitignore 覆盖缓存/工作目录/dist',
              '__pycache__/' in gi and '.workbuddy/' in gi and 'dist/' in gi))
    # ⑥ 现行文档不得引用 2.0.0 之前的旧章节号（现为 §0~§6）；CHANGELOG 属历史记录，回溯性引用合法，排除；
    #    回归项数随版本增长，硬编码必然过期 → 一律改为引用报告
    stale_sec, stale_num = [], []
    cand = [os.path.join(root, 'SKILL.md'), os.path.join(root, 'README.md'),
            os.path.join(root, 'CONTRIBUTING.md')]
    cand += glob.glob(os.path.join(root, 'references', '*.md'))
    for p in cand:
        if not os.path.exists(p):
            continue
        t = open(p, encoding='utf-8').read()
        bad = [s for s in re.findall(r'§(\d+)', t) if int(s) > 6]
        if bad:
            stale_sec.append('%s:%s' % (os.path.basename(p), ','.join(bad)))
        if re.search(r'(?<![\d.])\d{2,3}\s*/\s*\d{2,3}(?![\d.])', t):
            stale_num.append(os.path.basename(p))
    r.append(('现行文档无已废弃章节号(>§6)', not stale_sec, ';'.join(stale_sec)))
    r.append(('治理文件不硬编码回归项数(改引用报告)', not stale_num, ','.join(stale_num)))
    return r


def main():
    out = []
    EV = read_engine_version(ROOT)
    out.append('引擎版本：%s' % EV)
    files = sorted(glob.glob(os.path.join(ROOT, 'examples', '*.md')))
    total = ok = 0
    for f in files:
        name, rs = check(f)
        out.append('=== %s ===' % name)
        for item in rs:
            label, passed = item[0], item[1]
            note = item[2] if len(item) > 2 else ''
            total += 1
            ok += 1 if passed else 0
            out.append('  [%s] %s %s' % ('PASS' if passed else 'FAIL', label, note))
    out.append('')
    out.append('=== 结构化输出契约 references/output-schema.json ===')
    for item in check_schema(ROOT, EV):
        total += 1
        ok += 1 if item[1] else 0
        out.append('  [%s] %s %s' % ('PASS' if item[1] else 'FAIL', item[0], item[2] if len(item) > 2 else ''))
    out.append('')
    out.append('=== 引擎自身 SKILL.md ===')
    for item in check_skill_md(ROOT):
        total += 1
        ok += 1 if item[1] else 0
        out.append('  [%s] %s %s' % ('PASS' if item[1] else 'FAIL', item[0], item[2] if len(item) > 2 else ''))
    out.append('')
    out.append('=== md→JSON 派生一致性（Single Source） ===')
    for item in check_derived(ROOT, EV):
        total += 1
        ok += 1 if item[1] else 0
        out.append('  [%s] %s %s' % ('PASS' if item[1] else 'FAIL', item[0], item[2] if len(item) > 2 else ''))
    out.append('')
    out.append('=== Word 成品 md→docx（结构/版式/幂等/防漂移） ===')
    for item in check_docx(ROOT):
        total += 1
        ok += 1 if item[1] else 0
        out.append('  [%s] %s %s' % ('PASS' if item[1] else 'FAIL', item[0], item[2] if len(item) > 2 else ''))
    out.append('')
    out.append('=== 交付通道契约 md→docx CLI / PDF 降级 / 反篡改 ===')
    for item in check_delivery(ROOT):
        total += 1
        ok += 1 if item[1] else 0
        out.append('  [%s] %s %s' % ('PASS' if item[1] else 'FAIL', item[0], item[2] if len(item) > 2 else ''))
    out.append('')
    out.append('=== 双轨合规与仓库卫生 ===')
    for item in check_repo(ROOT, EV):
        total += 1
        ok += 1 if item[1] else 0
        out.append('  [%s] %s %s' % ('PASS' if item[1] else 'FAIL', item[0], item[2] if len(item) > 2 else ''))
    out.append('')
    # 元断言：同一分节内同一断言重复计入会让通过率虚高，必须为 0
    # （跨文件同名断言是合法的，故按分节重置去重集）
    seen, dup = set(), []
    for l in out:
        if l.startswith('==='):
            seen.clear()
            continue
        if l.startswith('  [PASS') or l.startswith('  [FAIL'):
            key = l.split(']', 1)[1].strip()
            if key in seen:
                dup.append(l.strip())
            seen.add(key)
    for d in dup[:5]:
        total += 1
        out.append('  [FAIL] 断言重复计入 %s' % d)
    out.append('')
    rate = ok * 100.0 / total if total else 0
    out.append('通过率：%d/%d = %.1f%%' % (ok, total, rate))
    open(report_path('regression_report.txt'), 'w', encoding='utf-8').write('\n'.join(out))


if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        import traceback
        open(report_path('regression_report.txt'), 'w',
             encoding='utf-8').write('ERROR:\n' + traceback.format_exc())
